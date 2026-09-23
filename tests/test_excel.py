"""
Excel hisobotlari sinovi.

Fayl yig'iladi, so'ng QAYTA O'QILIB har bir katakcha tekshiriladi —
"xato bermadi" yetarli emas, ma'lumot to'g'ri joyda turishi kerak.

Ishga tushirish:
    python tests/test_excel.py
"""

from __future__ import annotations

import asyncio
import io
import sys
from pathlib import Path

PROJECT_ROOT = Path(__file__).resolve().parent.parent
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

#  ⚠️ MUHIM: engine importidan OLDIN — sinov ALOHIDA bazani ishlatadi,
#  aks holda ishlaydigan botning ma'lumotlari o'chib ketadi.
from tests import _env  # noqa: F401,E402

from openpyxl import load_workbook  # noqa: E402

from core.security.permissions import Role  # noqa: E402
from infrastructure.database.engine import engine, session_factory  # noqa: E402
from modules.analytics.excel import AdminReportBuilder, TestReportBuilder  # noqa: E402
from modules.assessment.repository import AttemptRepository  # noqa: E402
from modules.assessment.service import AssessmentService  # noqa: E402
from modules.catalog.service import CatalogService  # noqa: E402
from modules.identity.models import User  # noqa: E402
from modules.identity.repository import UserRepository  # noqa: E402
from modules.registry import Base  # noqa: E402

KEY = "abcdabcdab"          # 10 savol

#  (telegram_id, ism, familiya, telefon, javoblari, kutilgan foiz)
STUDENTS = [
    (201, "Komronbek", "Risqulov", "+998901112233", "abcdabcdab", 100.0),
    (202, "Aziz", "Karimov", "+998901112244", "abcdabcdad", 90.0),
    (203, "Dilnoza", "Yusupova", None, "abcdabcdcc", 80.0),
    (204, "Sardor", "Toshev", "+998901112266", "abcd------", 40.0),
    (205, "Malika", "Ergasheva", "+998901112277", "dcbadcbadc", 0.0),
]

_failures: list[str] = []


def check(label: str, got, want) -> None:
    if got == want:
        print(f"  OK   {label}")
    else:
        print(f"  FAIL {label}\n         kutilgan: {want!r}\n         olingan : {got!r}")
        _failures.append(label)


async def make_user(session, telegram_id, first, last, role, phone=None) -> User:
    users = UserRepository(session)
    user, _ = await users.get_or_create(
        telegram_id, first_name=first, last_name=last,
        username=f"{first.lower()}_{last.lower()}",
    )
    user.role = role
    user.is_registered = True
    user.phone = phone
    await session.flush()
    return user


async def main() -> int:
    async with engine.begin() as connection:
        await connection.run_sync(Base.metadata.drop_all)
        await connection.run_sync(Base.metadata.create_all)

    # ==================================================================
    print("\n--- Tayyorgarlik ---")
    async with session_factory() as session:
        teacher = await make_user(
            session, 200, "Saxobiddin", "Karimov", Role.TEACHER.value
        )
        test = await CatalogService(session).create_from_one_line(
            teacher, f"Blok-3 Matematika+{KEY}"
        )
        number = test.number
        await session.commit()
        print(f"  OK   Test №{number} yaratildi, {test.questions_count} savol")

    for telegram_id, first, last, phone, answers, expected in STUDENTS:
        async with session_factory() as session:
            student = await make_user(
                session, telegram_id, first, last, Role.STUDENT.value, phone
            )
            catalog = CatalogService(session)
            found = await catalog.tests.get_by_number(number)
            result = await AssessmentService(session).submit(found, student, answers)
            check(f"{first} {last} natijasi", result.percentage, expected)
            await session.commit()

    # ==================================================================
    print("\n--- Test hisoboti yig'ildi ---")
    async with session_factory() as session:
        catalog = CatalogService(session)
        test = await catalog.tests.get_full(
            (await catalog.tests.get_by_number(number)).id
        )
        builder = TestReportBuilder(session)

        data, total = await builder.build(test)
        check("ishtirokchilar", total, 5)
        check("fayl bo'sh emas", len(data) > 3000, True)

        name = builder.filename(test)
        check("kengaytma", name.endswith(".xlsx"), True)
        check("nomda test raqami", f"(№{number})" in name, True)
        print(f"  ℹ️  {name}  ({len(data):,} bayt)")

        workbook = load_workbook(io.BytesIO(data))
        check("2 ta varaq", workbook.sheetnames, ["Natijalar", "Savollar tahlili"])

        # --- 1-varaq ---
        sheet = workbook["Natijalar"]
        check("sarlavha", sheet["A1"].value, "Blok-3 Matematika")

        headers = [sheet.cell(row=4, column=i).value for i in range(1, 16)]
        check("№ ustuni", headers[0], "№")
        check("ism ustuni", headers[1], "Familiya, Ism")
        check("telefon ustuni", headers[3], "Telefon")
        check("anti-cheat ustuni", headers[12], "Anti-cheat nazorati")
        check("javoblar ustuni", headers[13], "Javoblari")

        #  Reyting tartibida: 100% birinchi
        check("1-o'rin ismi", sheet.cell(row=5, column=2).value, "Komronbek Risqulov")
        check("1-o'rin username", sheet.cell(row=5, column=3).value, "@komronbek_risqulov")
        check("1-o'rin telefoni", sheet.cell(row=5, column=4).value, "+998901112233")
        check("1-o'rin foizi", sheet.cell(row=5, column=7).value, 1.0)
        check("1-o'rin anti-cheat", sheet.cell(row=5, column=13).value, "🟢 Halol (0)")
        check("1-o'rin javoblari", sheet.cell(row=5, column=14).value, "ABCDABCDAB")
        check("1-o'rin holati", sheet.cell(row=5, column=12).value, "O'tdi")

        check("oxirgi o'rin", sheet.cell(row=9, column=2).value, "Malika Ergasheva")
        check("oxirgi holati", sheet.cell(row=9, column=12).value, "O'tmadi")
        check("telefonsiz '—'", sheet.cell(row=7, column=4).value, "—")

        check("filtr yoqilgan", sheet.auto_filter.ref is not None, True)
        check("sarlavha qotirilgan", sheet.freeze_panes, "A5")

        # --- 2-varaq ---
        analysis = workbook["Savollar tahlili"]
        check("tahlil sarlavhasi", analysis["A1"].value, "Savollar bo'yicha tahlil")

        a_headers = [analysis.cell(row=4, column=i).value for i in range(1, 9)]
        check("savol ustuni", a_headers[0], "Savol №")
        check("muvaffaqiyat ustuni", a_headers[5], "Muvaffaqiyat %")
        check("chalg'ituvchi ustuni", a_headers[6], "Ko'p tanlangan xato")
        check("izoh ustuni", a_headers[7], "Izoh")

        #  1-savol: kalit 'a'. Javoblar: a,a,a,a,d -> 4 to'g'ri
        check("1-savol raqami", analysis.cell(row=5, column=1).value, 1)
        check("1-savol to'g'ri javobi", analysis.cell(row=5, column=2).value, "A")
        check("1-savolni 4 kishi topdi", analysis.cell(row=5, column=3).value, 4)
        check("1-savol muvaffaqiyati", analysis.cell(row=5, column=6).value, 0.8)

        #  1-savolda yagona xato 'd' bo'lgan
        check("1-savol chalg'ituvchisi", analysis.cell(row=5, column=7).value, "D (1)")

        #  9-savol: kalit 'a'. Javoblar: a,a,c,-,d -> 2 to'g'ri, 2 xato, 1 javobsiz
        check("9-savol to'g'ri", analysis.cell(row=13, column=3).value, 2)
        check("9-savol xato", analysis.cell(row=13, column=4).value, 2)
        check("9-savol javobsiz", analysis.cell(row=13, column=5).value, 1)

        #  Xatolar sochilib ketgan (C va D bittadan) — kalit shubhali emas
        check("9-savol izohi ogohlantirish emas", analysis.cell(row=13, column=8).value, "Qiyin")

        rows = sum(
            1 for r in range(5, 30)
            if analysis.cell(row=r, column=1).value is not None
        )
        check("10 ta savol qatori", rows, 10)

        workbook.close()

    # ==================================================================
    print("\n--- Admin: foydalanuvchilar hisoboti ---")
    async with session_factory() as session:
        users = await UserRepository(session).list_for_export()
        builder = AdminReportBuilder(session)

        data, total = await builder.build_users(users)
        check("foydalanuvchilar soni", total, 6)   # 1 o'qituvchi + 5 o'quvchi

        workbook = load_workbook(io.BytesIO(data))
        sheet = workbook["Foydalanuvchilar"]
        check("sarlavha", sheet["A1"].value, "Foydalanuvchilar")

        headers = [sheet.cell(row=4, column=i).value for i in range(1, 13)]
        check("rol ustuni", headers[5], "Rol")
        check("yechgan testlari ustuni", headers[7], "Yechgan testlari")

        #  Har bir qatorda ism bo'lishi kerak
        names = [sheet.cell(row=r, column=2).value for r in range(5, 11)]
        check("6 ta ism yozildi", sum(1 for n in names if n), 6)

        workbook.close()
        print(f"  ℹ️  {builder.users_filename()}  ({len(data):,} bayt)")

    # ==================================================================
    print("\n--- Admin: test yechganlar hisoboti ---")
    async with session_factory() as session:
        attempts = await AttemptRepository(session).list_recent_for_export()
        builder = AdminReportBuilder(session)

        data, total = await builder.build_solvers(attempts, period_label="Butun davr")
        check("topshiriqlar soni", total, 5)

        workbook = load_workbook(io.BytesIO(data))
        sheet = workbook["Test yechganlar"]
        check("sarlavha", sheet["A1"].value, "Test yechganlar")

        headers = [sheet.cell(row=4, column=i).value for i in range(1, 13)]
        check("test raqami ustuni", headers[4], "Test №")
        check("test nomi ustuni", headers[5], "Test nomi")

        check("test raqami yozildi", sheet.cell(row=5, column=5).value, number)
        check("test nomi yozildi", sheet.cell(row=5, column=6).value, "Blok-3 Matematika")

        workbook.close()
        print(f"  ℹ️  {builder.solvers_filename('Butun davr')}  ({len(data):,} bayt)")

    # ==================================================================
    print("\n--- Bo'sh hisobot xato bermaydi ---")
    async with session_factory() as session:
        builder = AdminReportBuilder(session)
        data, total = await builder.build_solvers([], period_label="Bugun")
        check("bo'sh hisobot yig'ildi", total, 0)
        check("fayl baribir tuzildi", len(data) > 1000, True)

    await engine.dispose()

    print()
    if _failures:
        print(f"  {len(_failures)} ta sinov muvaffaqiyatsiz:")
        for name in _failures:
            print(f"    - {name}")
        return 1

    print("  EXCEL HISOBOTLARI TO'G'RI")
    return 0


if __name__ == "__main__":
    raise SystemExit(asyncio.run(main()))
